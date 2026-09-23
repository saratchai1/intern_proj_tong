from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH = "รายงานสหกิจ.docx"
doc = Document(PATH)

def style_body(p, text):
    p.text = text
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def replace_placeholder(marker, paragraphs):
    target = None
    for t in doc.tables:
        if t.rows and marker in t.rows[0].cells[0].text:
            target = t
            break
    if target is None:
        print("WARN placeholder not found:", marker)
        return False
    tbl = target._element
    parent = tbl.getparent()
    idx = parent.index(tbl)
    for text in paragraphs:
        p = doc.add_paragraph()
        style_body(p, text)
        el = p._element
        el.getparent().remove(el)
        parent.insert(idx, el)
        idx += 1
    parent.remove(tbl)
    return True

chapter2 = {
"[เว้นไว้สำหรับเติมภายหลัง] อธิบายความสำคัญของ Drawing Standardization": [
"แบบวิศวกรรมเป็นเอกสารที่ใช้ถ่ายทอดข้อมูลจากผู้ออกแบบไปยังผู้ตรวจสอบ ผู้ควบคุมงาน ผู้รับเหมา และผู้เกี่ยวข้องอื่นในโครงการ จึงต้องอาศัยทั้งความถูกต้องของเนื้อหาทางวิศวกรรมและความสม่ำเสมอของรูปแบบเอกสาร การปรับมาตรฐาน Drawing ในโครงงานนี้หมายถึงการทำให้องค์ประกอบที่กำหนดไว้ล่วงหน้า เช่น ข้อมูลในกรอบแบบ รูปแบบตัวอักษร มาตราส่วน ขนาดกระดาษ และการตั้งค่า Plot มีรูปแบบที่สอดคล้องกันก่อนจัดพิมพ์หรือส่งมอบ โดยไม่ได้หมายถึงการแก้ไขเนื้อหาการออกแบบทางวิศวกรรมภายใน Model Space.",
"แนวคิดดังกล่าวสอดคล้องกับงานด้าน rule-based checking ซึ่งมองว่าข้อกำหนดที่นิยามเป็นเงื่อนไขได้ชัดเจนสามารถแปลงเป็นกฎสำหรับให้คอมพิวเตอร์ตรวจสอบได้ Murakami et al. (1995) ศึกษาการตรวจ CAD drawing ตาม drafting rules ขณะที่ Parfitt et al. (1993) เสนอกรอบระบบควบคุมคุณภาพอัตโนมัติสำหรับ working drawings แนวคิดทั้งสองสนับสนุนการเปลี่ยนมาตรฐานที่เป็นข้อความให้เป็นรายการตรวจสอบที่เครื่องสามารถประมวลผลได้.",
"สำหรับ AUTOSHEET การปรับมาตรฐานเริ่มจากการกำหนด Ground Truth และ Requirement R01–R27 ก่อนพัฒนาและประเมินระบบ โดยแยกข้อกำหนดเฉพาะโครงการ เช่น Cordia SHX, Text Height 0.0015 และ Plot Style ออกจากข้อกำหนดด้านความปลอดภัยของระบบ เช่น การรักษา Viewport Scale/View/Lock และการไม่แก้ profile ที่ระบบยังไม่รู้จัก."
],
"[เว้นไว้สำหรับเติมภายหลัง] อธิบาย Model Space, Paper Space, Layout และ Viewport": [
"AutoCAD แยกสภาพแวดล้อมหลักออกเป็น Model Space และ Paper Space โดย Model Space ใช้สำหรับสร้างและจัดการแบบจำลองหรือรูปทรงของงาน ส่วน Paper Space ใช้สำหรับจัดวางแบบลงบนแผ่นสำหรับการพิมพ์ แต่ละ Layout ใน Paper Space สามารถกำหนด Page Setup, ขนาดกระดาษ, เครื่องพิมพ์ และองค์ประกอบบนแผ่นแยกจากกันได้ การที่ไฟล์ DWG หนึ่งไฟล์มีหลาย Layout จึงทำให้การตรวจและแก้ไขแบบ Manual เกิดงานซ้ำจำนวนมาก.",
"Viewport เป็นหน้าต่างใน Paper Space ที่ใช้แสดง Model Space ด้วยมุมมองและมาตราส่วนที่กำหนด ดังนั้นการปรับแต่ง Paper Space ต้องระวังไม่ให้เปลี่ยน View, Viewport Scale หรือสถานะ Lock โดยไม่ตั้งใจ ใน AUTOSHEET ข้อกำหนด R20–R23 จึงกำหนดให้รักษาค่าเหล่านี้ไว้ เพื่อให้ระบบปรับมาตรฐานเอกสารโดยไม่เปลี่ยนสาระของแบบ."
],
"[เว้นไว้สำหรับเติมภายหลัง] อธิบายบทบาทของ Title Block": [
"Title Block ใช้ระบุข้อมูลสำคัญที่ทำให้ Drawing สามารถระบุตัวตนและติดตามสถานะได้ เช่น ชื่อแบบ รหัสแบบ ผู้จัดทำ ผู้ตรวจสอบ และข้อมูลโครงการ ISO 7200:2004 กำหนดแนวทางเกี่ยวกับ data fields ใน title blocks และ document headers เพื่อสนับสนุนการแลกเปลี่ยนเอกสารทางวิศวกรรม (International Organization for Standardization [ISO], 2004).",
"อย่างไรก็ตาม โครงสร้างและรูปแบบจริงของ Title Block ในแต่ละโครงการอาจแตกต่างกัน AUTOSHEET จึงใช้ Template และ Approved Drawing ของโครงการเป็น Ground Truth สำหรับค่ารายละเอียดเฉพาะ และมุ่งแก้เฉพาะ Drawing Title, Drawing Code และ Scale ที่อยู่ในขอบเขต.",
"ระหว่างพัฒนาเคยพบ Block ที่ถูกตีความว่าเป็น Title Box แต่จากการวิเคราะห์ Bounding Box พบว่าเป็น Full Sheet Frame ที่รวมเส้นกรอบไว้ด้วย ระบบรุ่นหลังจึงเปลี่ยนจากการลบและแทนที่เป็น Preserve Frame และแก้เฉพาะข้อมูลที่ต้องเปลี่ยน กรณีนี้สะท้อนความสำคัญของการตรวจโครงสร้างวัตถุจริงก่อน destructive edit."
],
"[เว้นไว้สำหรับเติมภายหลัง] อธิบาย Scale และ A-series drawing sheets": [
"มาตราส่วนเป็นความสัมพันธ์ระหว่างขนาดที่แสดงบน Drawing กับขนาดจริง ISO 5455:1979 ให้คำจำกัดความ ประเภท และรูปแบบการระบุมาตราส่วนสำหรับ technical drawings ส่วน ISO 5457:1999 กำหนดขนาดและ layout ของ drawing sheets รวมถึงชุดกระดาษ ISO-A (ISO, 1979; ISO, 1999).",
"ในโครงงานนี้ Requirement ด้าน Scale ใช้รูปแบบตาม Approved Drawing ของสถานประกอบการ การทดสอบพบว่า AUTOSHEET สามารถทำให้ค่าและรูปแบบข้อความ Scale ถูกต้อง แต่ R08 ซึ่งเกี่ยวข้องกับตำแหน่งเชิงเรขาคณิตของข้อความ A3/A1 ยังต้อง Manual Touch-up ในทุก Test Case ปัญหานี้แตกต่างจากกฎเชิงคุณสมบัติ เช่น Font หรือ Text Height เพราะต้องพิจารณาความสัมพันธ์เชิงตำแหน่งกับองค์ประกอบรอบข้าง.",
"สำหรับ lettering ใน technical documentation ISO 3098-1:2015 กำหนดข้อกำหนดทั่วไปของตัวอักษร แต่ Cordia SHX, Text Height 0.0015 และ Width Factor 1.0 ที่ใช้ใน AUTOSHEET เป็น Project-specific Ground Truth ไม่ใช่ค่าที่ ISO กำหนด (ISO, 2015)."
],
"[เว้นไว้สำหรับเติมภายหลัง] อธิบาย Plotter, paper, orientation, plot area, scale และ CTB": [
"การ Plot ใน AutoCAD ประกอบด้วยตัวแปรหลายรายการ เช่น Plotter/Printer, Paper Size, Orientation, Plot Area, Center Plot, Plot Scale, Plot Style และ Lineweight หากค่าดังกล่าวไม่สอดคล้องกันระหว่าง Layout ผลการพิมพ์อาจแตกต่างจากมาตรฐานแม้ Drawing บนหน้าจอจะดูถูกต้อง.",
"ใน AUTOSHEET การตั้งค่า Plot ถูกแปลงเป็น Requirement R09–R19 เช่น A3 Landscape, DWG To PDF.pc3, Plot Area แบบ Window, Center Plot, Custom Scale และ SIPHYA_LT_R1.ctb และตรวจค่าหลัง Apply เพื่อยืนยันว่า Layout พร้อมสำหรับ Plot ตาม Ground Truth.",
"ระหว่างพัฒนาเคยพบกรณี findfile ไม่พบไฟล์ CTB ทั้งที่ AutoCAD สามารถเลือก Plot Style ดังกล่าวได้ ระบบรุ่นหลังจึงเปลี่ยนเป็นอ่านและตรวจค่าจาก Layout/AutoCAD โดยตรง แทนการใช้การมีอยู่ของไฟล์ใน search path เป็นเงื่อนไขเพียงอย่างเดียว."
],
"[เว้นไว้สำหรับเติมภายหลัง] อธิบาย AutoLISP และการใช้เพื่อ automate": [
"AutoLISP เป็นภาษาสำหรับปรับแต่งและขยายความสามารถของ AutoCAD โดย Autodesk อธิบายว่าสามารถใช้สร้างคำสั่งและโปรแกรมตั้งแต่งานง่ายที่เรียกใช้คำสั่งซ้ำ ไปจนถึงการจัดการวัตถุและคุณสมบัติของ Drawing (Autodesk, 2026a, 2026b).",
"ความสามารถในการเข้าถึงวัตถุ ตรวจชนิดและคุณสมบัติ วนทำงานกับหลาย Layout และกำหนดเงื่อนไข ทำให้ AutoLISP เหมาะกับงานซ้ำ เช่น การเปลี่ยน Drawing Code, ปรับข้อความ, ตั้ง Page Setup และตรวจผลหลังดำเนินการ.",
"AUTOSHEET เริ่มจาก LISP เฉพาะงานหลายคำสั่งก่อนพัฒนาเป็น workflow ที่มี Scan, Apply และ Re-scan โดยยังคง Final QA โดยผู้ใช้ ระบบจึงมีเป้าหมายลด repetitive work มากกว่าทดแทนผู้ปฏิบัติงานทั้งหมด."
],
"[เว้นไว้สำหรับเติมภายหลัง] สังเคราะห์แนวคิด rule-based checking": [
"Rule-based checking คือแนวคิดที่แปลงข้อกำหนดหรือมาตรฐานให้เป็นเงื่อนไขที่เครื่องสามารถตรวจสอบได้ งานของ Murakami et al. (1995) แสดงแนวคิดดังกล่าวกับ CAD drawing ส่วน Eastman et al. (2009) และ Solihin and Eastman (2015) อธิบายกรอบและความซับซ้อนของกฎในระบบตรวจแบบอัตโนมัติ แม้งานสองชิ้นหลังอยู่ในบริบท BIM แต่สามารถใช้สนับสนุนหลักการว่ากฎที่ชัดเจนและตรวจได้จาก property เหมาะต่อ automation มากกว่ากฎที่ต้องอาศัยความสัมพันธ์เชิงพื้นที่และบริบท.",
"กรอบแนวคิดของ AUTOSHEET จึงสรุปเป็น Standard/Requirement → Machine-checkable Rule → Automatic Check/Adjustment → Human Validation โดย Requirement ที่เป็น property-based เช่น Font, Text Height, Drawing Code และ Plot Setting สามารถตรวจและแก้ได้โดยตรง ขณะที่ R08 Scale Alignment ต้องพิจารณา geometry และยังคงใช้ Human Validation."
],
"[เว้นไว้สำหรับเติมภายหลัง] เขียนสังเคราะห์งานวิจัยด้าน rule-based CAD checking": [
"Murakami et al. (1995) และ Parfitt et al. (1993) เป็นงานพื้นฐานที่สนับสนุนแนวคิดการตรวจคุณภาพ Drawing จากกฎ ขณะที่ Scheibel et al. (2021) แสดงว่าข้อมูลจาก engineering drawings สามารถถูกสกัดและนำไปใช้สนับสนุน quality control ได้ งานเหล่านี้ชี้ว่าการจัดโครงสร้าง requirement ให้เครื่องประมวลผลเป็นแนวทางที่นำไปใช้ได้จริง.",
"Rica et al. (2020) ศึกษาการลด human effort ใน engineering drawing validation และสะท้อนแนวคิดว่า automation สามารถช่วยลดภาระผู้ตรวจได้โดยยังคง human validation ในจุดที่จำเป็น ซึ่งสอดคล้องกับ AUTOSHEET ที่ยังต้อง Final Visual QA และ Manual Touch-up สำหรับ R08.",
"Sofias et al. (2025) แสดงการใช้ CAD API automation กับ recurrent engineering tasks ใน workflow จริงและประเมินผ่าน case studies แนวทางนี้ใกล้เคียงกับวิธีประเมิน AUTOSHEET อย่างไรก็ตาม รายงานนี้ไม่ใช้ผลเชิงตัวเลขของงานภายนอกเป็นผลของ AUTOSHEET แต่ใช้ T01–T05 เป็นหลักฐานเชิงประจักษ์ของระบบเอง."
]
}

for k,v in chapter2.items():
    replace_placeholder(k,v)

chapter5 = {
"[เว้นไว้สำหรับเติมภายหลัง] ตอบวัตถุประสงค์ทั้ง 4 ข้อ": [
"ผลการดำเนินโครงงานตอบวัตถุประสงค์ทั้ง 4 ข้อ ได้แก่ (1) ได้ศึกษากระบวนการจัดเตรียมและตรวจสอบ Drawing ก่อน Plot จากงานจริงที่มีประมาณ 300–500 ไฟล์ DWG และมากกว่า 1,000 แผ่น (2) ได้จำแนกงานซ้ำที่สามารถกำหนดเป็นกฎและนำไป automate ได้ (3) ได้พัฒนาต้นแบบ AUTOSHEET และ workflow ที่ใช้ AutoLISP เพื่อปรับมาตรฐานและตรวจสอบแบบ และ (4) ได้ประเมินประสิทธิภาพเทียบกับกระบวนการ Manual ด้วย Test Case T01–T05.",
"ในการประเมิน 5 ไฟล์ รวม 15 Layout เวลารวมของ Manual Workflow เท่ากับ 29 นาที 11.52 วินาที ขณะที่ Auto-assisted Workflow ซึ่งรวม AUTOSHEET, A3TITLEBOXALL และ Manual Touch-up ใช้เวลา 15 นาที 51.80 วินาที หรือลดลงประมาณ 45.7% ในชุดทดสอบนี้ ก่อน Touch-up ระบบผ่าน Requirement เต็มรูปแบบ 26 จาก 27 ข้อ และหลัง Final QA ทุก Test Case อยู่ในสถานะ Ready to Plot."
],
"[เว้นไว้สำหรับเติมภายหลัง] เชื่อมผลการทดลองกับงานวิจัยในบทที่ 2": [
"ผลการทดสอบแสดงให้เห็นว่าข้อกำหนดที่เป็น property-based หรือ explicit rule เช่น Font, Text Height, Drawing Code, Paper และ Plot Setting สามารถตรวจและปรับได้อย่างสม่ำเสมอกว่าข้อกำหนดที่พึ่งพาความสัมพันธ์เชิงพื้นที่ ผลนี้สอดคล้องกับแนวคิด rule-based checking ของ Murakami et al. (1995) และกรอบความซับซ้อนของกฎที่ Solihin and Eastman (2015) อธิบายไว้.",
"R08 Scale Alignment เป็น Requirement เดียวที่ต้อง Manual Touch-up ใน T01–T05 แม้ค่า Scale และรูปแบบข้อความจะถูกต้องแล้ว สะท้อนว่าการตรวจค่าข้อมูลกับการตัดสินตำแหน่งเชิง geometry เป็นปัญหาคนละระดับ การคง Human Validation ไว้จึงเหมาะสมกับต้นแบบนี้และสอดคล้องกับ Rica et al. (2020) ซึ่งเสนอ automation เพื่อลดภาระมนุษย์โดยไม่จำเป็นต้องตัดผู้เชี่ยวชาญออกจากกระบวนการทั้งหมด.",
"การประเมินจาก workflow จริงและการจับเวลาจน Ready to Plot มีลักษณะเป็น case-study evaluation คล้ายแนวทางของ Sofias et al. (2025) ที่ประเมิน CAD API automation กับ recurrent engineering tasks อย่างไรก็ตาม ผล 45.7% ของรายงานนี้ใช้เฉพาะชุดทดสอบ AUTOSHEET และไม่ควรนำไปตีความเป็นค่าทั่วไปของ CAD automation."
],
"[เว้นไว้สำหรับเติมภายหลัง] สรุปข้อจำกัดจากผลจริง": [
"การประเมินมีข้อจำกัดหลายประการ ชุดทดสอบประกอบด้วย 5 ไฟล์ รวม 15 Layout ซึ่งเหมาะสำหรับการประเมินต้นแบบในบริบทสหกิจศึกษา แต่ยังไม่ครอบคลุม Drawing profile ทุกประเภท และ Drawing ที่ใช้มาจากบริบทงานที่เกี่ยวข้องกัน จึงไม่ควรสรุปว่าอัตราการลดเวลา 45.7% จะเกิดขึ้นกับ Drawing ทุกชนิด.",
"ผู้พัฒนาระบบเป็นผู้ดำเนินการทดลอง Manual และ Auto-assisted Workflow เอง ความคุ้นเคยกับ Drawing และเครื่องมืออาจมีผลต่อเวลา และแต่ละ Test Case วัดเวลาเพียงหนึ่งรอบหลักโดยไม่มี repeated measurement หลายรอบ ดังนั้นผลด้านเวลาควรตีความเป็น case-study evaluation ไม่ใช่ benchmark สากล.",
"ระบบยังต้องใช้ลำดับ AUTOSHEET → A3TITLEBOXALL เพื่อให้ R25 ผ่าน และ R08 Scale Alignment ยังต้อง Manual Touch-up ในทุก Test Case Final Visual QA จึงยังจำเป็น นอกจากนี้ รายงานปกปิดหรือครอบตัดรายละเอียด Drawing ที่ไม่จำเป็นต่อการอธิบายผลเพื่อรักษาความลับของสถานประกอบการ."
],
"[เว้นไว้สำหรับเติมภายหลัง] สรุปประโยชน์ที่พิสูจน์ได้จากผลทดลอง": [
"ประโยชน์ที่พิสูจน์ได้คือการลดเวลาจนถึง Ready to Plot จาก 29:11.52 นาที เหลือ 15:51.80 นาที หรือประมาณ 45.7% ในชุดทดสอบ โดยตัวเลขดังกล่าวรวมเวลาของ automation และ Manual Touch-up แล้ว จึงสะท้อนเวลาที่ผู้ใช้ต้องใช้จริงมากกว่าการรายงานเฉพาะ runtime ของสคริปต์.",
"ด้านคุณภาพ ระบบทำให้ Requirement ที่เป็นกฎชัดเจนถูกตรวจและปรับอย่างสม่ำเสมอ โดยผ่านเต็มรูปแบบ 26 จาก 27 ข้อก่อน Touch-up และหลัง Final QA ทุกกรณี Ready to Plot ประโยชน์จึงรวมทั้งความเร็วและการลดโอกาสข้ามขั้นตอนซ้ำ เช่น Plot Setting หรือการปรับข้อมูลหลาย Layout."
],
"[เว้นไว้สำหรับเติมภายหลัง] เสนอแนวทาง เช่น configurable profiles": [
"การพัฒนาระยะถัดไปควรให้ความสำคัญกับ R08 โดยใช้ geometry และ bounding box ของข้อความกับองค์ประกอบรอบข้าง เพื่อจัดวาง A3/A1 ได้ยืดหยุ่นขึ้น พร้อมคง preview หรือ diagnostic ก่อน destructive edit.",
"ควรแยก Project Standard ออกจาก code หลักเป็น configurable profiles เช่น Font, Text Height, Title Block profile, CTB, Paper และกฎเฉพาะโครงการ เพื่อลด hard-code และรองรับ Drawing variant อื่นได้ง่ายขึ้น.",
"ควรเพิ่ม audit trail หลังรัน เช่น Requirement ที่ผ่าน ไม่ผ่าน หรือถูกข้าม พร้อมเหตุผล และเพิ่ม regression test จาก Drawing ที่ได้รับอนุญาตให้ใช้ทดสอบ รวมถึงขยาย Evaluation Set และให้ผู้ใช้มากกว่าหนึ่งคนทดลองเพื่อประเมิน usability และ generalizability."
]
}
for k,v in chapter5.items():
    replace_placeholder(k,v)

# Remove bibliography placeholder
for t in list(doc.tables):
    if t.rows and "[เว้นไว้สำหรับเติมภายหลัง] เพิ่มเติม Autodesk Official Documentation" in t.rows[0].cells[0].text:
        t._element.getparent().remove(t._element)

# Append references after Sofias reference if not already present
refs = [
"Autodesk. (2026a). AutoLISP Developer's Guide. AutoCAD 2026 Help. https://help.autodesk.com/cloudhelp/2026/ENU/AutoCAD-AutoLISP/files/GUID-265AADB3-FB89-4D34-AA9D-6ADF70FF7D4B.htm",
"Autodesk. (2026b). Introduction (AutoLISP). AutoCAD 2026 Help. https://help.autodesk.com/cloudhelp/2026/CSY/AutoCAD-AutoLISP/files/GUID-A0E9D801-8BE9-4BF1-85E8-3807E15F3B71.htm",
"International Organization for Standardization. (1979). Technical drawings — Scales (ISO 5455:1979). https://www.iso.org/standard/11500.html",
"International Organization for Standardization. (1999). Technical product documentation — Sizes and layout of drawing sheets (ISO 5457:1999). https://www.iso.org/standard/29017.html",
"International Organization for Standardization. (2004). Technical product documentation — Data fields in title blocks and document headers (ISO 7200:2004). https://www.iso.org/standard/35446.html",
"International Organization for Standardization. (2015). Technical product documentation — Lettering — Part 1: General requirements (ISO 3098-1:2015). https://www.iso.org/standard/65679.html",
"เอกสารมาตรฐานและแบบอ้างอิงของโครงการ. (2569). Template, Approved Drawing และ Plot Setting ที่ใช้เป็น Ground Truth ในโครงงาน AUTOSHEET [เอกสารภายในสถานประกอบการ ไม่เผยแพร่]."
]
existing = {p.text.strip() for p in doc.paragraphs}
anchor = next((p for p in doc.paragraphs if p.text.startswith("Sofias, K., Kanetaki")), None)
if anchor:
    cur = anchor
    for ref in refs:
        if ref in existing:
            continue
        np = doc.add_paragraph(ref)
        np.alignment = WD_ALIGN_PARAGRAPH.LEFT
        np.paragraph_format.left_indent = Inches(0.5)
        np.paragraph_format.first_line_indent = Inches(-0.5)
        np.paragraph_format.space_after = Pt(3)
        el = np._element
        el.getparent().remove(el)
        cur._element.addnext(el)
        cur = np

# Fix page header page-number fields
for s in doc.sections:
    h = s.header
    if s is doc.sections[0] and not any((p.text or "").strip() for p in h.paragraphs):
        continue
    for p in list(h.paragraphs):
        p._element.getparent().remove(p._element)
    p = h.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ["ascii","hAnsi","eastAsia","cs"]:
        rf.set(qn("w:"+a), "Arial")
    rpr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)

# Restart numbered lists in 1.2, 1.3, 1.5 if still using List Number
paras = doc.paragraphs
for start,end in [("1.2 วัตถุประสงค์","1.3 ขอบเขตการศึกษา"),("1.3 ขอบเขตการศึกษา","1.4 วิธีการศึกษา"),("1.5 ประโยชน์ที่คาดว่าจะได้รับ","บทที่ 2")]:
    try:
        si = next(i for i,p in enumerate(paras) if p.text.strip() == start)
        ei = next(i for i,p in enumerate(paras) if i > si and p.text.strip() == end)
    except StopIteration:
        continue
    n = 1
    for p in paras[si+1:ei]:
        if p.style.name == "List Number":
            txt = p.text
            p.style = doc.styles["Normal"]
            p.text = f"{n}. {txt}"
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            n += 1

# Keep quantitative statements and references left aligned to avoid stretched digits
for p in doc.paragraphs:
    if any(k in p.text for k in ["29 นาที 11.52","15 นาที 51.80","29:11.52","15:51.80","45.7%","26 จาก 27","27 จาก 27"]):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if p.text.startswith(("Murakami,","Parfitt,","Rica,","Scheibel,","Eastman,","Solihin,","Sofias,","Autodesk.","International Organization","เอกสารมาตรฐาน")):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)

doc.save(PATH)
print("Updated", PATH)
