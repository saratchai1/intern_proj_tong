from docx import Document
from docx.shared import Inches

PATH="รายงานสหกิจ.docx"
doc=Document(PATH)

def find_exact(text):
    for p in doc.paragraphs:
        if p.text.strip()==text:
            return p
    return None

def find_start(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None

def insert_after(anchor,text,style="Normal"):
    p=doc.add_paragraph(style=style)
    p.text=text
    el=p._element
    el.getparent().remove(el)
    anchor._element.addnext(el)
    return p

def clear_between(a,b):
    parent=a._element.getparent()
    cur=a._element.getnext()
    while cur is not None and cur is not b._element:
        nxt=cur.getnext()
        if cur.tag.endswith("}p"):
            parent.remove(cur)
        cur=nxt

def replace_section(start_text,end_text,items):
    s=find_exact(start_text); e=find_exact(end_text)
    if not s or not e:
        print("WARN section not found",start_text,end_text); return
    clear_between(s,e)
    a=s
    for style,text in items:
        a=insert_after(a,text,style)

def set_start(prefix,text):
    p=find_start(prefix)
    if not p:
        print("WARN para not found",prefix); return
    p.text=text

def set_single_table(prefix,text):
    for t in doc.tables:
        if len(t.rows)==1 and len(t.rows[0].cells)==1 and t.cell(0,0).text.strip().startswith(prefix):
            t.cell(0,0).text=text; return
    print("WARN abstract table not found",prefix)

def set_table(header_first,rows,qualifier=None):
    for t in doc.tables:
        if not t.rows or not t.rows[0].cells: continue
        if t.rows[0].cells[0].text.strip()!=header_first: continue
        if qualifier and qualifier not in [c.text.strip() for c in t.rows[0].cells]: continue
        while len(t.rows)<len(rows): t.add_row()
        while len(t.rows)>len(rows): t._tbl.remove(t.rows[-1]._tr)
        for ri,row in enumerate(rows):
            for ci,val in enumerate(row):
                t.cell(ri,ci).text=str(val)
        return
    print("WARN table not found",header_first,qualifier)

TH_ABS="""โครงงานนี้มีวัตถุประสงค์เพื่อพัฒนาต้นแบบระบบอัตโนมัติสำหรับช่วยปรับมาตรฐานและตรวจสอบคุณภาพแบบวิศวกรรมก่อนการจัดพิมพ์ใน AutoCAD ด้วยภาษา AutoLISP จากการปฏิบัติงานจริงพบว่าชุด Drawing ที่ได้รับมีขนาดรวมประมาณ 9 GB ประกอบด้วยไฟล์ DWG ประมาณ 300–500 ไฟล์ และมีแบบมากกว่า 1,000 แผ่น โดยกระบวนการเดิมต้องตรวจและแก้ไของค์ประกอบหลายรายการซ้ำในแต่ละ Layout เช่น Title Block, Drawing Code, Font, Scale, Paper Setup และ Plot Setting จึงพัฒนาเครื่องมือ AutoLISP แบบเฉพาะงานและต่อยอดเป็น AUTOSHEET ร่วมกับคำสั่ง A3TITLEBOXALL ภายใต้ workflow ที่กำหนดไว้ หลังการทดสอบนำร่อง T01 ได้ตรึงวิธีประเมินและใช้ T02–T05 เป็นชุดประเมินหลักจำนวน 4 ไฟล์ รวม 14 Layout ผลการประเมินพบว่า Manual Workflow ใช้เวลารวม 27 นาที 6.10 วินาที ขณะที่ Auto-assisted Workflow ซึ่งรวม AUTOSHEET, A3TITLEBOXALL, Manual Touch-up และ Final QA ใช้เวลารวม 14 นาที 57.34 วินาที ลดเวลารวมประมาณ 44.8% โดยค่าการลดเวลารายกรณีอยู่ในช่วง 33.4–55.1% และมีค่ามัธยฐาน 46.6% ด้านคุณภาพ กระบวนการอัตโนมัติผ่านข้อกำหนดเต็มรูปแบบ 26 จาก 27 ข้อในทุกกรณีประเมินหลัก ส่วน R08 Scale Alignment อยู่ในระดับ Partially Compliant เนื่องจากค่ามาตราส่วนและรูปแบบข้อความถูกต้องแต่ตำแหน่ง A3/A1 ยังต้องปรับด้วยผู้ใช้งาน หลัง Manual Touch-up ทุกกรณีผ่านข้อกำหนดที่ applicable และอยู่ในสถานะ Ready to Plot ผลดังกล่าวชี้ว่าระบบสามารถลดงานซ้ำและเพิ่มความสม่ำเสมอของ Drawing ได้ในบริบทที่ทดสอบ โดยยังคง Human Validation สำหรับกฎที่ขึ้นกับ geometry และบริบทของแบบ"""
EN_ABS="""This cooperative education project developed an AutoLISP-based prototype to support engineering drawing standardization and pre-plot quality control in AutoCAD. The assigned drawing set comprised approximately 9 GB of data, 300–500 DWG files, and more than 1,000 sheets. The original manual workflow required repetitive checking and editing across layouts, including title blocks, drawing codes, fonts, scales, paper setup, and plot settings. Task-specific AutoLISP utilities were therefore developed iteratively and integrated into an AUTOSHEET-assisted workflow together with A3TITLEBOXALL. After pilot test T01 was used to refine and freeze the evaluation protocol, T02–T05 were used as the main evaluation set, covering 4 DWG files and 14 layouts. The manual workflow required 27 min 6.10 s in total, whereas the auto-assisted workflow, including AUTOSHEET, A3TITLEBOXALL, manual touch-up, and final QA, required 14 min 57.34 s, corresponding to a pooled time reduction of approximately 44.8%. Case-level reductions ranged from 33.4% to 55.1%, with a median of 46.6%. Before manual touch-up, the automated workflow achieved full compliance with 26 of 27 applicable requirements in every evaluation case. R08 Scale Alignment remained partially compliant because scale values and formatting were correct while A3/A1 text positioning still required manual adjustment. After touch-up, all applicable requirements were satisfied and every case reached Ready-to-Plot status. The results indicate that rule-based CAD automation can reduce repetitive work and improve consistency in the evaluated context while retaining human validation for geometry- and context-dependent conditions."""
set_single_table("โครงงานนี้มีวัตถุประสงค์",TH_ABS)
set_single_table("This cooperative education project",EN_ABS)

set_start("การศึกษานี้เริ่มจากการสำรวจกระบวนการจัดเตรียม Drawing","""การศึกษานี้เริ่มจากการสำรวจกระบวนการจัดเตรียม Drawing ที่ดำเนินการด้วยวิธี Manual และวิเคราะห์ขั้นตอนที่เกิดซ้ำ จากนั้นพัฒนา AutoLISP ขนาดเล็กเพื่อช่วยตรวจสอบและแก้ปัญหาเฉพาะงาน ก่อนจะปรับปรุงคำสั่งเหล่านั้นจากผลการทดลองกับ Drawing จริงและรวบรวมฟังก์ชันที่เกี่ยวข้องเป็น AUTOSHEET ต่อมาได้กำหนดข้อกำหนดของ Drawing จาก Template, Approved Drawing และข้อกำหนดของโครงการเป็น Ground Truth โดยใช้ T01 เป็น Pilot สำหรับกำหนดและตรึง Protocol จากนั้นจึงใช้ T02–T05 เป็นชุดประเมินหลักภายใต้ Workflow คงที่ AUTOSHEET → A3TITLEBOXALL → Auto-only QA → Manual Touch-up/Final QA และเปรียบเทียบกับกระบวนการ Manual""")
set_start("หลักการด้านความปลอดภัยที่ใช้ร่วมกับ Workflow ได้แก่","""หลักการด้านความปลอดภัยที่ใช้ร่วมกับ Workflow ได้แก่ การไม่เปลี่ยน Viewport Scale, View และ Lock โดยไม่จำเป็น; การ Preserve Full Sheet Frame ใน Profile ที่ตรวจพบว่าเป็นกรอบทั้งแผ่น; การ Skip และแจ้ง Warning เมื่อ Profile ไม่อยู่ในรูปแบบที่รองรับ; และการไม่บันทึกไฟล์โดยอัตโนมัติ เพื่อให้ผู้ใช้งานตรวจสอบผลก่อนตัดสินใจ Save ทั้ง Pilot T01 และชุดประเมินหลัก T02–T05 ใช้นิยามเดียวกันว่า “กระบวนการอัตโนมัติ” หมายถึง AUTOSHEET → A3TITLEBOXALL ตามลำดับนี้เท่านั้น""")

sec39=[
("Heading 2","3.9.1 Pilot Test และการ Freeze วิธีประเมิน"),
("Normal","การประเมินเริ่มจาก T01 ซึ่งกำหนดให้เป็น Pilot Test เพื่อทดสอบความเป็นไปได้ของ protocol ตรวจจุดผิดพลาดของ workflow และกำหนดลำดับคำสั่งสุดท้าย โดยใช้ไฟล์ต้นฉบับเดียวกันสร้างสำเนาสำหรับ Manual และ Auto-assisted Workflow เมื่อ Pilot พบประเด็น R25 และยืนยันลำดับ AUTOSHEET → A3TITLEBOXALL แล้ว วิธีประเมินจึงถูก Freeze ก่อนเริ่มชุดประเมินหลัก ดังนั้น T01 ใช้เป็นหลักฐานเชิงพัฒนาและตัวอย่าง Before/Auto/Final แต่ไม่รวมอยู่ในสถิติหลักของประสิทธิภาพระบบ."),
("Heading 2","3.9.2 การเลือกชุดประเมินหลัก"),
("Normal","ชุดประเมินหลักประกอบด้วย T02–T05 จำนวน 4 ไฟล์ รวม 14 Layout คัดเลือกแบบ purposive sampling หลังจาก Freeze workflow เพื่อให้ครอบคลุมไฟล์ที่มีจำนวน Layout แตกต่างกัน 2–5 Layout และมีองค์ประกอบเป้าหมายของระบบ เช่น Title Block, Scale, Plot Setting และ Cleanup อยู่ในขอบเขตที่ประเมิน การคัดเลือกนี้มีวัตถุประสงค์เพื่อประเมินต้นแบบในบริบทงานจริง ไม่ได้ออกแบบให้เป็นตัวอย่างสุ่มที่แทน Drawing ทั้ง 300–500 ไฟล์ในเชิงสถิติ และไม่มีการเปลี่ยน Algorithm ตามผล T02–T05 ระหว่างการประเมิน."),
("Heading 2","3.9.3 Timing Protocol และเกณฑ์ Ready to Plot"),
("Normal","การทดลองดำเนินโดยผู้พัฒนาระบบคนเดียวในสภาพแวดล้อมการทำงานเดียวกัน เพื่อให้เงื่อนไขระหว่าง Manual และ Auto-assisted ใกล้เคียงกัน T_manual เริ่มจับเมื่อเริ่มดำเนินการกับสำเนาไฟล์ต้นฉบับและหยุดเมื่อผู้ปฏิบัติงานตรวจครบทุก Layout และพิจารณาว่า Ready to Plot; T_auto เริ่มก่อนเรียก AUTOSHEET และหยุดเมื่อ A3TITLEBOXALL ทำงานเสร็จ; T_touchup เริ่มหลังตรวจผล Auto-only และหยุดเมื่อแก้จุดที่ยังไม่ผ่านพร้อมทำ Final Visual QA ครบทุก Layout."),
("Normal","ในแต่ละ case ลำดับการทดลองเป็น Manual ก่อน Auto-assisted จึงอาจเกิด order/learning effect เพราะผู้ทดลองคุ้นเคยกับ Drawing มากขึ้นในช่วงหลัง ประเด็นนี้ถูกระบุเป็นข้อจำกัดและไม่มีการใช้ repeated measurement เพื่อประมาณความแปรปรวนเชิงสถิติ."),
("Heading 2","3.9.4 QA Protocol และนิยาม R08"),
("Normal","หลัง AUTOSHEET → A3TITLEBOXALL จะประเมิน R01–R27 โดยยังไม่แก้ด้วยมือ ใช้สถานะ PASS, PARTIAL, WARN, FAIL และ N/A สำหรับ R08 Scale Alignment กำหนด operational definition ว่า Scale A1 ต้องอยู่ใต้ Scale A3 และจัดแนวตาม arrangement ของ Approved Drawing/Template โดยไม่เหลื่อมหรือแยกออกจาก Scale group อย่างเห็นได้ชัด เนื่องจาก Ground Truth ด้านตำแหน่งเป็นเกณฑ์เชิงภาพตามแบบอ้างอิงและไม่ได้กำหนด tolerance เป็นมิลลิเมตร จึงรายงาน R08 เป็น PARTIAL เมื่อค่าและรูปแบบข้อความถูกต้องแต่ตำแหน่งยังต้องปรับด้วยผู้ใช้."),
("Heading 2","3.9.5 ตัวชี้วัดและการวิเคราะห์"),
("Normal","ตัวชี้วัดหลักประกอบด้วย (1) เวลา Manual, Auto และ Touch-up (2) Full-compliance status ของ R01–R27 (3) Automation Coverage และ (4) Residual Manual Work โดยจำนวน User Operations ใช้เป็นข้อมูลประกอบเฉพาะ Pilot เนื่องจากไม่ได้บันทึกอย่างสม่ำเสมอทุก case."),
("Normal","Effective Time Reduction (%) = [(T_manual - T_auto+touchup) / T_manual] × 100 โดย T_auto+touchup = T_auto + T_touchup การสรุปผลหลักใช้ descriptive statistics ของ T02–T05 ได้แก่ pooled time reduction, median และ range ของ time reduction ราย case ไม่ทำ inferential statistical test เนื่องจากมีเพียง 4 evaluation cases และแต่ละ case วัดเวลาหนึ่งรอบ."),
("Normal","ด้านความลับ ภาพ Drawing ในรายงานใช้การ crop/redact และใช้รหัส Evidence ภายในแทนชื่อไฟล์จริงเมื่อรายละเอียดไม่จำเป็นต่อการอธิบายผล Full evidence ถูกเก็บแยกใน Evidence Pack ภายในเพื่อให้ตรวจสอบย้อนหลังได้ภายใต้สิทธิ์ของสถานประกอบการ.")
]
replace_section("3.9 การออกแบบการทดลองและการเก็บข้อมูล","3.10 หลักเกณฑ์การแก้ Code หลัง Pilot และการ Freeze รุ่นประเมินผล",sec39)

s=find_exact("3.10 หลักเกณฑ์การแก้ Code หลัง Pilot และการ Freeze รุ่นประเมินผล"); e=find_exact("บทที่ 4")
if s and e:
    clear_between(s,e)
    a=insert_after(s,"หลัง Pilot T01 กำหนดให้แก้เฉพาะข้อผิดพลาดที่เป็น Blocker หรือมีความเสี่ยงต่อความถูกต้องของ Drawing เช่น ลบ Object ผิด ทำให้ Scale value ผิด Plot Setting ผิด ทำให้กรอบแบบหรือ Viewport เสียหาย หรือทำให้ไฟล์ไม่สามารถใช้ต่อได้ ส่วน R08 ซึ่งเป็นปัญหาการจัดตำแหน่งข้อความ Scale A3/A1 โดยที่ค่าและรูปแบบมาตราส่วนยังถูกต้อง ถูกจัดเป็นข้อจำกัดแบบ Partially Automated และวัดเป็น Residual Manual Work แทนการแก้ Algorithm ต่อระหว่างชุดประเมินหลัก T02–T05.")
    insert_after(a,"เมื่อ T01 รอบ Final ไม่พบ Blocker อื่น จึง Freeze วิธีประเมินเป็น Evaluation Workflow V1.0 โดยใช้ลำดับ AUTOSHEET → A3TITLEBOXALL → Auto-only QA → Manual Touch-up/Final QA จากนั้นใช้ลำดับเดียวกันกับ T02–T05 และไม่เปลี่ยน Algorithm ระหว่างชุดประเมินหลัก เพื่อแยกข้อมูลที่ใช้ปรับ protocol ออกจากข้อมูลที่ใช้สรุปประสิทธิภาพระบบ.")

updates={
"รุ่นที่ใช้ในการประเมินผลถูกกำหนดเป็น Workflow คงที่":"Workflow สำหรับการประเมินถูกตรึงหลัง Pilot T01 โดยประกอบด้วยการรัน AUTOSHEET ก่อน แล้วรัน A3TITLEBOXALL ต่อทันที จากนั้นตรวจ R01–R27 ก่อนเข้าสู่ Manual Touch-up และ Final QA ลำดับเดียวกันนี้ถูกใช้กับชุดประเมินหลัก T02–T05 และห้ามสลับคำสั่ง เนื่องจาก A3TITLEBOXALL เป็นขั้นตอนที่ทำให้การจัดการ Title Block และ Source/Workspace Path ตาม R25 ครบถ้วนใน Workflow ที่ทดสอบจริง",
"เมื่อรวมทั้ง 5 กรณี":"T01 แสดงไว้เป็น Pilot เพื่อให้เห็นกระบวนการพัฒนา protocol แต่สถิติหลักคำนวณจากชุดประเมิน T02–T05 จำนวน 4 ไฟล์ รวม 14 Layout กระบวนการ Manual ใช้เวลารวม 27 นาที 6.10 วินาที ขณะที่ Auto-assisted Workflow ใช้เวลารวม 14 นาที 57.34 วินาที คิดเป็น pooled time reduction 44.8% และใช้เวลาประมาณ 55.2% ของ Manual Workflow การลดเวลาราย case อยู่ในช่วง 33.4–55.1% และมีค่ามัธยฐาน 46.6% เวลาของคำสั่งอัตโนมัติ T02–T05 รวม 1 นาทีพอดี ส่วนเวลาที่เหลือส่วนใหญ่เป็น Manual Touch-up และ Final QA.",
"ตาราง 4.3 ผล Standard Compliance":"ตาราง 4.3 ผล Full-compliance ของชุดประเมินหลัก T02–T05",
"Residual Manual Work ที่เกิดซ้ำในทุกกรณี":"Residual Manual Work ที่เกิดซ้ำในชุดประเมินหลัก T02–T05 คือการจัดตำแหน่งข้อความมาตราส่วน A3/A1 ตาม R08 และการตรวจด้วยสายตาก่อน Plot ไม่พบข้อกำหนดอื่นที่ต้องแก้ด้วยมืออย่างสม่ำเสมอในชุดประเมินหลัก โดย Pilot T01 พบรูปแบบเดียวกัน.",
"ไม่พบ Destructive Failure ในชุดประเมิน T01–T05":"ไม่พบ Destructive Failure ในชุดประเมินหลัก T02–T05 เช่น การลบกรอบแบบผิด การเปลี่ยน Viewport Scale/View/Lock โดยไม่ตั้งใจ หรือ Plot Setting ผิดจาก Ground Truth ข้อจำกัดที่พบซ้ำคือ R08 Scale Alignment ซึ่งถูกจัดเป็น PARTIAL ไม่ใช่ Failure ของค่า Scale เนื่องจาก Scale content และรูปแบบ A3/A1 ถูกต้อง แต่ตำแหน่งข้อความต้องเลื่อนด้วยมือ.",
"นอกจากนี้ การลบ Source/Workspace Path (R25)":"นอกจากนี้ การลบ Source/Workspace Path (R25) จะผ่านเมื่อใช้ Workflow ตามลำดับ AUTOSHEET → A3TITLEBOXALL เท่านั้น จึงกำหนดลำดับดังกล่าวเป็นส่วนหนึ่งของ Evaluation Workflow และใช้เหมือนกันใน T02–T05 หลังจาก Freeze ที่ Pilot T01.",
"การประเมิน T01–T05 ครอบคลุม 5 ไฟล์":"เมื่อแยก T01 เป็น Pilot ชุดประเมินหลัก T02–T05 ครอบคลุม 4 ไฟล์ รวม 14 Layout กระบวนการ Manual ใช้เวลารวม 27:06.10 นาที ขณะที่ Auto-assisted Workflow ซึ่งประกอบด้วย AUTOSHEET → A3TITLEBOXALL → Manual Touch-up/Final QA ใช้เวลารวม 14:57.34 นาที ลดเวลารวมได้ประมาณ 44.8% ในชุดประเมินนี้ โดยค่าลดเวลาราย case มีมัธยฐาน 46.6% และช่วง 33.4–55.1%."
}
for k,v in updates.items(): set_start(k,v)

set_table("Function",[
["Function","Fully Automated","Partially Automated","Manual Required","หมายเหตุ"],
["Title Block / Frame","Yes","No","No","ผ่านใน Workflow AUTOSHEET → A3TITLEBOXALL"],
["Font / Text Height","Yes","No","No","R01–R03 ผ่านทุกกรณี"],
["Drawing Code","Yes","No","No","R04 ผ่านทุกกรณี"],
["Paper / Plot Setting","Yes","No","No","R09–R19 ผ่านทุกกรณี"],
["Scale Content","Yes","No","No","R06–R07 ผ่านทุกกรณี"],
["Scale Alignment","No","Yes","Yes","R08 ต้องเลื่อน A3/A1 ด้วยมือใน T02–T05; Pilot T01 พบรูปแบบเดียวกัน"],
["Cleanup","Yes","No","No","R24–R25 ผ่าน; R25 สำเร็จด้วย A3TITLEBOXALL"],
["Final Layout QA","No","No","Yes","กำหนดให้ตรวจทุก Layout ก่อน Plot โดยผู้ใช้งาน"],
["Viewport / Safety","Yes","No","No","R20–R23 และ R26–R27 ผ่านทุกกรณี"]
])
set_table("Test",[
["Test","Layouts","Manual","Auto","Touch-up","Auto+Touch","Reduce %"],
["T01 (Pilot)","1","02:05.42","00:11.38","00:43.08","00:54.46","56.6*"],
["T02","4","09:11.24","00:11.56","04:48.34","04:59.90","45.6"],
["T03","3","06:21.30","00:12.54","04:01.52","04:14.06","33.4"],
["T04","2","04:29.30","00:12.36","01:48.54","02:00.90","55.1"],
["T05","5","07:04.26","00:23.54","03:18.94","03:42.48","47.6"],
["Evaluation T02–T05","14","27:06.10","01:00.00","13:57.34","14:57.34","44.8"]
],qualifier="Manual")
set_table("Requirement Group",[
["Requirement Group","ผลหลัง Auto workflow (T02–T05)","ผลหลัง Touch-up","ข้อสังเกต"],
["R01–R07","PASS 4/4","PASS 4/4","Font/Title/Code/Scale content ผ่านตาม Ground Truth"],
["R08 Scale Alignment","PARTIAL 4/4","PASS 4/4","ต้องเลื่อนข้อความ A3/A1 ด้วยมือทุกกรณี"],
["R09–R27","PASS 4/4","PASS 4/4","รวม R25 ซึ่งสำเร็จหลัง A3TITLEBOXALL"],
["Full-compliance ratio","26/27 = 96.3%","27/27 = 100%","รายงานเป็น full compliance; R08 แยกเป็น PARTIAL"]
])
set_table("Test",[
["Test","Layouts","Touch-up Time","Residual Manual Work","Final Status"],
["T01 (Pilot)","1","00:43.08","จัดตำแหน่ง Scale A3/A1 + Final visual QA","Ready to Plot"],
["T02","4","04:48.34","จัดตำแหน่ง Scale A3/A1 ในหลาย Layout + Final visual QA","Ready to Plot"],
["T03","3","04:01.52","จัดตำแหน่ง Scale A3/A1 + Final visual QA","Ready to Plot"],
["T04","2","01:48.54","จัดตำแหน่ง Scale A3/A1 + Final visual QA","Ready to Plot"],
["T05","5","03:18.94","จัดตำแหน่ง Scale A3/A1 ในหลาย Layout + Final visual QA","Ready to Plot"]
],qualifier="Touch-up Time")

sections={
"5.1 สรุปผลตามวัตถุประสงค์":[
("Normal","โครงงานนี้กำหนดวัตถุประสงค์ไว้ 4 ข้อ และผลการดำเนินงานสามารถสรุปตามวัตถุประสงค์ได้ดังนี้"),
("Normal","วัตถุประสงค์ข้อที่ 1 เพื่อศึกษากระบวนการจัดเตรียมและตรวจสอบแบบวิศวกรรมใน AutoCAD ก่อนการจัดพิมพ์หรือส่งมอบ ผลจากการศึกษากระบวนการทำงานจริงพบว่าการเตรียม Drawing ต้องตรวจองค์ประกอบหลายรายการในแต่ละ Layout ได้แก่ Title Block, Drawing Title, Drawing Code, Font, Scale, Paper, Plot Setting และข้อมูลที่ไม่ต้องการ กระบวนการเดิมมีงานซ้ำจำนวนมากโดยเฉพาะเมื่อไฟล์มีหลาย Layout ซึ่งนำไปสู่การกำหนด Manual Workflow และ Ground Truth สำหรับการประเมินในโครงงาน."),
("Normal","วัตถุประสงค์ข้อที่ 2 เพื่อวิเคราะห์ขั้นตอนการทำงานซ้ำที่สามารถประยุกต์ใช้ระบบอัตโนมัติ ผลการวิเคราะห์แยกงานออกเป็นงานที่สามารถนิยามเป็นกฎได้ชัดเจน งานที่ automate ได้บางส่วน และงานที่ยังต้องอาศัยการพิจารณาของผู้ใช้ ข้อกำหนด R01–R27 ถูกใช้เป็น Requirement Master โดยพบว่ากฎด้าน property และ configuration เช่น Font, Text Height, Drawing Code, Paper และ Plot Setting เหมาะต่อการทำ automation มากกว่ากฎที่ขึ้นกับตำแหน่งเชิงเรขาคณิต."),
("Normal","วัตถุประสงค์ข้อที่ 3 เพื่อพัฒนาต้นแบบ AUTOSHEET ด้วย AutoLISP สำหรับช่วยปรับมาตรฐานและตรวจสอบ Drawing ระบบถูกพัฒนาจาก LISP เฉพาะงานหลายรุ่นและปรับปรุงจากผลทดสอบจริงจนเป็น workflow ที่ใช้ AUTOSHEET ตามด้วย A3TITLEBOXALL จากนั้นจึงตรวจ Auto-only QA และ Manual Touch-up/Final QA หลัง Freeze workflow ไม่พบ destructive failure ในชุดประเมินหลัก T02–T05 และหลัง Manual Touch-up ทุก case อยู่ในสถานะ Ready to Plot."),
("Normal","วัตถุประสงค์ข้อที่ 4 เพื่อประเมินประสิทธิภาพเมื่อเทียบกับกระบวนการ Manual หลังใช้ T01 เป็น Pilot และ Freeze protocol ชุดประเมินหลัก T02–T05 จำนวน 4 ไฟล์ รวม 14 Layout พบว่า Manual ใช้เวลารวม 27 นาที 6.10 วินาที ขณะที่ Auto-assisted Workflow ใช้เวลารวม 14 นาที 57.34 วินาที ลดเวลารวมประมาณ 44.8% ในชุดประเมินนี้ ด้าน Full-compliance ระบบผ่านข้อกำหนดเต็มรูปแบบ 26 จาก 27 ข้อหลัง automated workflow ในทุก case โดย R08 Scale Alignment เป็น PARTIAL และเมื่อทำ Manual Touch-up แล้วทุกกรณีผ่าน 27/27 applicable requirements และพร้อม Plot.")
],
"5.2 อภิปรายผล":[
("Normal","ผลการทดลองแสดงให้เห็นว่าการแปลง Requirement ที่มีเงื่อนไขชัดเจนเป็นกฎของระบบสามารถช่วยให้การตรวจและปรับ Drawing มีความสม่ำเสมอมากขึ้น ผลนี้สอดคล้องกับแนวคิดของ Murakami et al. (1995) ที่ใช้ drafting rules สำหรับ computerized checking และสอดคล้องในเชิงแนวคิดกับกรอบ rule-based checking ของ Eastman et al. (2009) ซึ่งแยกกระบวนการเตรียมข้อมูล กฎ การประมวลผล และผลการตรวจออกเป็นขั้นตอน."),
("Normal","R08 เกิดซ้ำในชุดประเมินหลัก T02–T05 ทั้ง 4/4 cases และพบรูปแบบเดียวกันใน Pilot T01 ช่วยให้เห็นความแตกต่างระหว่าง rule เชิง property กับ rule เชิง geometry อย่างชัดเจน Font, Text Height, Drawing Code และ Plot Setting สามารถกำหนด target value และตรวจสอบตรง ๆ ได้ ขณะที่ Scale Alignment ต้องพิจารณาความสัมพันธ์เชิงตำแหน่งของข้อความกับองค์ประกอบรอบข้าง จึงมีความซับซ้อนมากกว่า ประเด็นนี้สอดคล้องในเชิงแนวคิดกับ Solihin and Eastman (2015) ที่ชี้ว่ากฎสำหรับ automated checking มีระดับความซับซ้อนและ prerequisite แตกต่างกัน แม้งานดังกล่าวศึกษา BIM ไม่ใช่ 2D AutoCAD."),
("Normal","ผลด้านเวลาควรพิจารณาจาก Auto-assisted Workflow ซึ่งรวม Manual Touch-up ไม่ใช่เฉพาะ runtime ของ LISP เพราะเป้าหมายเชิงปฏิบัติคือเวลาไปถึง Ready to Plot แนวคิดนี้สอดคล้องกับ Rica et al. (2020) ที่คงบทบาทของวิศวกรไว้ในขั้น human validation และมุ่งลด human effort ในส่วนที่ระบบอัตโนมัติสามารถช่วยได้ ในทำนองเดียวกัน AUTOSHEET ไม่ได้มีเป้าหมายแทนที่ผู้ปฏิบัติงาน แต่ลดงานซ้ำก่อน Final QA."),
("Normal","Sofias et al. (2025) แสดงการใช้ CAD API automation กับ recurrent engineering tasks ใน workflow จริงและประเมินผลด้วย case studies ซึ่งเป็นแนวทางที่ใกล้เคียงกับการประเมิน AUTOSHEET ในระดับวิธีวิจัย อย่างไรก็ตาม ผลเชิงตัวเลขของแต่ละงานไม่สามารถนำมาเปรียบเทียบตรง ๆ ได้ เนื่องจากโปรแกรม งาน และชุดข้อมูลแตกต่างกัน รายงานนี้จึงใช้ผล T02–T05 เป็นหลักฐานหลักของประสิทธิภาพ AUTOSHEET หลัง Freeze และใช้ T01 เป็น Pilot evidence ส่วนงานวิจัยภายนอกใช้สนับสนุนกรอบแนวคิดเท่านั้น.")
],
"5.3 ข้อจำกัดของ AUTOSHEET":[
("Normal","การประเมินของโครงงานมีข้อจำกัดที่ต้องคำนึงถึงในการตีความผล ประการแรก T01 เป็น Pilot และไม่รวมในสถิติหลัก ส่วนชุดประเมินหลักประกอบด้วย T02–T05 จำนวน 4 ไฟล์ รวม 14 Layout ซึ่งเหมาะสำหรับ case-study evaluation ของต้นแบบ แต่ยังไม่ครอบคลุม Drawing profile ทุกประเภท การเลือกตัวอย่างเป็น purposive sampling จากบริบทงานที่เกี่ยวข้องกัน จึงไม่ควรสรุปว่าอัตราการลดเวลา 44.8% จะเกิดขึ้นกับ Drawing ทุกชนิดหรือเป็นตัวแทนของ DWG ทั้งหมดในองค์กร."),
("Normal","ประการที่สอง ผู้พัฒนาระบบเป็นผู้ดำเนินการทดลอง Manual และ Auto-assisted Workflow เอง และใช้ลำดับ Manual ก่อน Auto-assisted ในแต่ละ case จึงมีโอกาสเกิด operator bias และ order/learning effect จากความคุ้นเคยกับ Drawing นอกจากนี้แต่ละ case วัดเวลาเพียงหนึ่งรอบหลักและไม่มี repeated measurement จึงไม่สามารถประมาณความแปรปรวนหรือทำ inferential statistical test ได้ ผลด้านเวลาจึงควรตีความเป็น descriptive case-study evidence ไม่ใช่ benchmark สากล."),
("Normal","ประการที่สาม ระบบยังต้องใช้ลำดับคำสั่ง AUTOSHEET → A3TITLEBOXALL เพื่อให้ Requirement R25 ผ่าน และ R08 Scale Alignment ยังต้อง Manual Touch-up ในทุก case ของชุดประเมินหลัก Final Visual QA จึงยังจำเป็นเพื่อป้องกันความผิดปกติที่ rule-based check อาจไม่ครอบคลุม ข้อจำกัดเหล่านี้สะท้อนว่า AUTOSHEET เป็นต้นแบบระบบสนับสนุนผู้ปฏิบัติงาน ไม่ใช่ระบบ Fully Autonomous."),
("Normal","ประการที่สี่ การยืนยัน Ready to Plot และการตรวจ QA ในการทดลองนี้ดำเนินการโดยผู้พัฒนาระบบเอง ยังไม่มี independent reviewer ตรวจผลทุก case จึงควรเพิ่มการตรวจรับโดยพนักงานที่ปรึกษาหรือผู้ใช้งานอื่นในงานต่อเนื่อง. ด้านการนำเสนอข้อมูล รายงานฉบับนี้จำกัดการเปิดเผยรายละเอียด Drawing ตามข้อกำหนดด้านความลับของสถานประกอบการ ภาพตัวอย่างจึงถูกครอบตัดหรือปกปิดข้อมูลที่ไม่จำเป็นต่อการอธิบายผล ขณะที่หลักฐานเต็มถูกเก็บใน Evidence Pack ภายใน.")
],
"5.4 ประโยชน์ต่อสถานประกอบการ":[
("Normal","ประโยชน์ที่พิสูจน์ได้จากชุดประเมินหลัก T02–T05 คือการลดเวลาจนถึง Ready to Plot จาก 27:06.10 นาที เหลือ 14:57.34 นาที หรือประมาณ 44.8% โดยการคำนวณรวมเวลาของ AUTOSHEET, A3TITLEBOXALL และ Manual Touch-up แล้ว จึงสะท้อนเวลาที่ผู้ใช้ต้องใช้จริงมากกว่าการรายงานเฉพาะ runtime ของ automation. ค่าลดเวลาราย case มีมัธยฐาน 46.6% และช่วง 33.4–55.1%."),
("Normal","ด้านคุณภาพ กระบวนการอัตโนมัติทำให้ Requirement ที่เป็นกฎชัดเจนถูกตรวจและปรับอย่างสม่ำเสมอในชุดประเมินหลัก T02–T05 โดยผ่านเต็มรูปแบบ 26 จาก 27 ข้อก่อน touch-up ในทุก case และหลัง Final QA ทุกกรณีอยู่ในสถานะ Ready to Plot ประโยชน์สำคัญจึงไม่ใช่เพียงความเร็ว แต่รวมถึงการลดโอกาสที่ผู้ใช้จะข้ามขั้นตอนที่เกิดซ้ำ เช่น การตั้งค่า Plot หรือการปรับรูปแบบข้อมูลในหลาย Layout."),
("Normal","ในเชิงการทำงานของสถานประกอบการ AUTOSHEET ยังช่วยแยกงานออกเป็นส่วนที่ระบบทำได้อัตโนมัติและส่วนที่ต้องให้ผู้ใช้ใช้เวลาในการพิจารณา ส่งผลให้ผู้ปฏิบัติงานสามารถให้ความสำคัญกับ Final QA และจุดที่ต้องใช้ judgment มากกว่าการทำ operation เดิมซ้ำหลายครั้ง.")
],
"5.5 แนวทางการพัฒนาต่อ":[
("Normal","แนวทางพัฒนาระยะถัดไปควรให้ความสำคัญกับ R08 Scale Alignment โดยพัฒนาการจัดวางข้อความจาก rule ที่อิงค่าคงที่ไปสู่การตรวจ geometry และ bounding box ของข้อความและองค์ประกอบรอบข้าง เพื่อรองรับกรณีที่ตำแหน่ง Title/Scale แตกต่างกันระหว่าง Drawing โดยควรคงกลไก preview หรือ diagnostic ก่อน destructive edit เพื่อรักษาความปลอดภัย."),
("Normal","ควรแยก Drawing profile และ Project Standard ออกจาก code หลักให้เป็น configuration ที่ปรับได้ เช่น Font, Text Height, Title Block profile, CTB, Paper และกฎที่ใช้กับแต่ละโครงการ วิธีนี้จะลดการ hard-code และทำให้ระบบนำไปใช้กับ Drawing variant อื่นได้ง่ายขึ้น."),
("Normal","อีกแนวทางหนึ่งคือเพิ่ม audit trail และรายงานผลหลังรัน เช่น Requirement ใดผ่าน ไม่ผ่าน หรือถูกข้าม พร้อมเหตุผล และบันทึกเวลาหรือจำนวนรายการที่แก้ เพื่อให้ผู้ตรวจสอบสามารถย้อนดูผลการทำงานได้โดยไม่ต้องพึ่งความจำของผู้ใช้ นอกจากนี้อาจพัฒนาชุด regression test จาก Drawing ที่ได้รับอนุญาตให้ใช้ทดสอบ เพื่อป้องกันการแก้ code ในอนาคตแล้วกระทบ Requirement ที่เคยผ่าน."),
("Normal","การพัฒนาในอนาคตควรเพิ่ม Evaluation Set ที่มีความหลากหลายทั้งจำนวน Layout, Title Block profile และประเภท Drawing รวมถึงทำ repeated timing, สลับลำดับการทดลองหรือใช้ผู้ทดสอบหลายคน และให้ independent reviewer เช่น พนักงานที่ปรึกษาตรวจสถานะ Ready to Plot เพื่อประเมินทั้ง repeatability, usability และ generalizability ของระบบได้ชัดเจนขึ้น.")
]
}
order=["5.1 สรุปผลตามวัตถุประสงค์","5.2 อภิปรายผล","5.3 ข้อจำกัดของ AUTOSHEET","5.4 ประโยชน์ต่อสถานประกอบการ","5.5 แนวทางการพัฒนาต่อ"]
for i,h in enumerate(order):
    replace_section(h, order[i+1] if i+1<len(order) else "เอกสารอ้างอิง", sections[h])

for p in doc.paragraphs:
    if p.text.startswith("Autodesk. (2013b, July 2). About plot style tables."):
        url=p.text.split("About plot style tables.",1)[1].strip()
        p._element.clear_content()
        p.add_run("Autodesk. (2013b, April 10). ")
        r=p.add_run("About plot style tables"); r.italic=True
        p.add_run(". "+url)
        p.paragraph_format.left_indent=Inches(0.5); p.paragraph_format.first_line_indent=Inches(-0.5)
    if p.text.startswith("Autodesk. (n.d.). AutoLISP developer's guide (AutoLISP)."):
        url=p.text.split("from ",1)[1]
        p._element.clear_content()
        p.add_run("Autodesk. (n.d.). ")
        r=p.add_run("AutoLISP developer's guide (AutoLISP)"); r.italic=True
        p.add_run(". Retrieved September 23, 2026, from "+url)
        p.paragraph_format.left_indent=Inches(0.5); p.paragraph_format.first_line_indent=Inches(-0.5)

for p in list(doc.paragraphs):
    if p.text.startswith("หมายเหตุเกี่ยวกับแหล่งข้อมูลภายใน: Approved Drawing"):
        p._element.getparent().remove(p._element)

doc.save(PATH)
print("Updated",PATH)
