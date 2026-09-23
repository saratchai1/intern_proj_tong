from docx import Document

PATH="รายงานสหกิจ.docx"
doc=Document(PATH)

TH_ABS="""โครงงานนี้มีวัตถุประสงค์เพื่อพัฒนาต้นแบบระบบอัตโนมัติสำหรับช่วยปรับมาตรฐานและตรวจสอบคุณภาพแบบวิศวกรรมก่อนการจัดพิมพ์ใน AutoCAD ด้วยภาษา AutoLISP จากการปฏิบัติงานจริงพบว่าชุด Drawing ที่ได้รับมีขนาดรวมประมาณ 9 GB ประกอบด้วยไฟล์ DWG ประมาณ 300–500 ไฟล์ และมีแบบมากกว่า 1,000 แผ่น โดยกระบวนการเดิมต้องตรวจและแก้ไของค์ประกอบหลายรายการซ้ำในแต่ละ Layout เช่น Title Block, Drawing Code, Font, Scale, Paper Setup และ Plot Setting จึงพัฒนาเครื่องมือ AutoLISP แบบเฉพาะงานและต่อยอดเป็น AUTOSHEET ร่วมกับคำสั่ง A3TITLEBOXALL ภายใต้ workflow ที่กำหนดไว้ หลังการทดสอบนำร่อง T01 ได้ตรึงวิธีประเมินและใช้ T02–T05 เป็นชุดประเมินหลักจำนวน 4 ไฟล์ รวม 14 Layout ผลการประเมินพบว่า Manual Workflow ใช้เวลารวม 27 นาที 6.10 วินาที ขณะที่ Auto-assisted Workflow ซึ่งรวม AUTOSHEET, A3TITLEBOXALL, Manual Touch-up และ Final QA ใช้เวลารวม 14 นาที 57.34 วินาที ลดเวลารวมประมาณ 44.8% โดยค่าการลดเวลารายกรณีอยู่ในช่วง 33.4–55.1% และมีค่ามัธยฐาน 46.6% ด้านคุณภาพ กระบวนการอัตโนมัติผ่านข้อกำหนดเต็มรูปแบบ 26 จาก 27 ข้อในทุกกรณีประเมินหลัก ส่วน R08 Scale Alignment อยู่ในระดับ Partially Compliant เนื่องจากค่ามาตราส่วนและรูปแบบข้อความถูกต้องแต่ตำแหน่ง A3/A1 ยังต้องปรับด้วยผู้ใช้งาน หลัง Manual Touch-up ทุกกรณีผ่านข้อกำหนดที่ applicable และอยู่ในสถานะ Ready to Plot ผลดังกล่าวชี้ว่าระบบสามารถลดงานซ้ำและเพิ่มความสม่ำเสมอของ Drawing ได้ในบริบทที่ทดสอบ โดยยังคง Human Validation สำหรับกฎที่ขึ้นกับ geometry และบริบทของแบบ"""

EN_ABS="""This cooperative education project developed an AutoLISP-based prototype to support engineering drawing standardization and pre-plot quality control in AutoCAD. The assigned drawing set comprised approximately 9 GB of data, 300–500 DWG files, and more than 1,000 sheets. The original manual workflow required repetitive checking and editing across layouts, including title blocks, drawing codes, fonts, scales, paper setup, and plot settings. Task-specific AutoLISP utilities were therefore developed iteratively and integrated into an AUTOSHEET-assisted workflow together with A3TITLEBOXALL. After pilot test T01 was used to refine and freeze the evaluation protocol, T02–T05 were used as the main evaluation set, covering 4 DWG files and 14 layouts. The manual workflow required 27 min 6.10 s in total, whereas the auto-assisted workflow, including AUTOSHEET, A3TITLEBOXALL, manual touch-up, and final QA, required 14 min 57.34 s, corresponding to a pooled time reduction of approximately 44.8%. Case-level reductions ranged from 33.4% to 55.1%, with a median of 46.6%. Before manual touch-up, the automated workflow achieved full compliance with 26 of 27 applicable requirements in every evaluation case. R08 Scale Alignment remained partially compliant because scale values and formatting were correct while A3/A1 text positioning still required manual adjustment. After touch-up, all applicable requirements were satisfied and every case reached Ready-to-Plot status. The results indicate that rule-based CAD automation can reduce repetitive work and improve consistency in the evaluated context while retaining human validation for geometry- and context-dependent conditions."""

# Front-matter abstract placeholder tables are tables 3 and 4 in the current report.
for t in doc.tables:
    txt=t.cell(0,0).text.strip() if t.rows and t.rows[0].cells else ""
    if txt.startswith("[เว้นไว้สำหรับเติมภายหลัง] เขียนหลังได้ผลการทดลองจริง"):
        t.cell(0,0).text=TH_ABS
    elif txt.startswith("[เว้นไว้สำหรับเติมภายหลัง] Translate from the finalized Thai abstract"):
        t.cell(0,0).text=EN_ABS

# Finalize body captions and List of Tables terminology.
for p in doc.paragraphs:
    if p.text.strip()=="ตาราง 2.1 Literature Matrix เบื้องต้น":
        p.text="ตาราง 2.1 การสังเคราะห์งานวิจัยที่เกี่ยวข้องกับ AUTOSHEET"
    elif p.text.strip()=="ตาราง 4.3 ผล Standard Compliance ของชุดประเมิน T01–T05":
        p.text="ตาราง 4.3 ผล Full-compliance ของชุดประเมินหลัก T02–T05"
    elif p.text.strip()=="ตาราง 4.3 ผล Standard Compliance":
        p.text="ตาราง 4.3 ผล Full-compliance ของชุดประเมินหลัก T02–T05"

# Update List of Tables entries in front matter.
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            c.text=c.text.replace("ตาราง 2.1 Literature Matrix เบื้องต้น","ตาราง 2.1 การสังเคราะห์งานวิจัยที่เกี่ยวข้องกับ AUTOSHEET")
            c.text=c.text.replace("ตาราง 4.3 ผล Standard Compliance","ตาราง 4.3 ผล Full-compliance ของชุดประเมินหลัก T02–T05")

doc.save(PATH)
print("Front matter and table titles finalized")
